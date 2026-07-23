#!/usr/bin/env python3
"""Convert LaTeX algorithm pseudocode to Word document format."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell, **kwargs):
    """Set cell border."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            for attr, val in kwargs[edge].items():
                element.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def parse_latex_algorithm(tex_content):
    """Parse LaTeX algorithm content and extract structured data."""
    lines = tex_content.split('\n')
    
    # Extract caption
    caption_match = re.search(r'\\caption\{(.+?)\}', tex_content)
    caption = caption_match.group(1) if caption_match else "Algorithm"
    
    # Extract Require
    require_match = re.search(r'\\Require\s+(.+?)(?=\\Ensure|$)', tex_content, re.DOTALL)
    require_text = ""
    if require_match:
        require_raw = require_match.group(1).strip()
        # Clean up LaTeX commands
        require_text = clean_latex(require_raw)
    
    # Extract Ensure
    ensure_match = re.search(r'\\Ensure\s+(.+?)(?=\\Statex|\\If|\\While|\\State|\\ForAll|\\EndIf|\\EndWhile|\\end\{algorithmic\})', tex_content, re.DOTALL)
    ensure_text = ""
    if ensure_match:
        ensure_raw = ensure_match.group(1).strip()
        ensure_text = clean_latex(ensure_raw)
    
    # Extract algorithm lines
    algorithm_lines = []
    in_algorithm = False
    line_counter = 0
    
    for line in lines:
        stripped = line.strip()
        
        if '\\begin{algorithmic}' in stripped:
            in_algorithm = True
            continue
        if '\\end{algorithmic}' in stripped:
            in_algorithm = False
            continue
        if not in_algorithm:
            continue
            
        # Skip Require/Ensure (already extracted)
        if '\\Require' in stripped or '\\Ensure' in stripped:
            continue
            
        # Handle Statex (section comments)
        statex_match = re.search(r'\\Statex\s+\\textit\{(.+?)\}', stripped)
        if statex_match:
            comment = statex_match.group(1)
            algorithm_lines.append(('comment', comment))
            continue
            
        # Handle regular comments
        if stripped.startswith('//'):
            algorithm_lines.append(('comment', stripped))
            continue
            
        # Handle State lines
        state_match = re.search(r'\\State\s+(.+)', stripped)
        if state_match:
            content = state_match.group(1)
            line_counter += 1
            algorithm_lines.append(('state', line_counter, clean_latex(content)))
            continue
            
        # Handle If
        if_match = re.search(r'\\If\{(.+?)\}', stripped)
        if if_match:
            line_counter += 1
            condition = clean_latex(if_match.group(1))
            algorithm_lines.append(('if', line_counter, condition))
            continue
            
        # Handle Else
        if '\\Else' in stripped:
            algorithm_lines.append(('else', None))
            continue
            
        # Handle ElsIf
        elsif_match = re.search(r'\\ElsIf\{(.+?)\}', stripped)
        if elsif_match:
            line_counter += 1
            condition = clean_latex(elsif_match.group(1))
            algorithm_lines.append(('elsif', line_counter, condition))
            continue
            
        # Handle While
        while_match = re.search(r'\\While\{(.+?)\}', stripped)
        if while_match:
            line_counter += 1
            condition = clean_latex(while_match.group(1))
            algorithm_lines.append(('while', line_counter, condition))
            continue
            
        # Handle EndIf
        if '\\EndIf' in stripped:
            algorithm_lines.append(('endif', None))
            continue
            
        # Handle EndWhile
        if '\\EndWhile' in stripped:
            algorithm_lines.append(('endwhile', None))
            continue
            
        # Handle ForAll
        forall_match = re.search(r'\\ForAll\{(.+?)\}', stripped)
        if forall_match:
            line_counter += 1
            condition = clean_latex(forall_match.group(1))
            algorithm_lines.append(('forall', line_counter, condition))
            continue
            
        # Handle EndFor
        if '\\EndFor' in stripped:
            algorithm_lines.append(('endfor', None))
            continue
            
        # Handle Return
        return_match = re.search(r'\\Return\s*(.*)', stripped)
        if return_match:
            line_counter += 1
            content = return_match.group(1).strip()
            if content:
                content = clean_latex(content)
            algorithm_lines.append(('return', line_counter, content))
            continue
    
    return caption, require_text, ensure_text, algorithm_lines


def clean_latex(text):
    """Clean LaTeX formatting commands."""
    # Remove \mathit{...}
    text = re.sub(r'\\mathit\{(.+?)\}', r'\1', text)
    # Remove \textsc{...}
    text = re.sub(r'\\textsc\{(.+?)\}', r'\1', text)
    # Remove \text{...}
    text = re.sub(r'\\text\{(.+?)\}', r'\1', text)
    # Remove \textit{...}
    text = re.sub(r'\\textit\{(.+?)\}', r'\1', text)
    # Replace \gets with ←
    text = text.replace('\\gets', '←')
    # Replace \Vert with ||
    text = text.replace('\\Vert', '∥')
    # Replace \leq with ≤
    text = text.replace('\\leq', '≤')
    # Replace \geq with ≥
    text = text.replace('\\geq', '≥')
    # Replace \neq with ≠
    text = text.replace('\\neq', '≠')
    # Replace \in with ∈
    text = text.replace('\\in', '∈')
    # Replace \exists with ∃
    text = text.replace('\\exists', '∃')
    # Replace \neg with ¬
    text = text.replace('\\neg', '¬')
    # Replace \perp with ⊥
    text = text.replace('\\perp', '⊥')
    # Replace \emptyset with ∅
    text = text.replace('\\emptyset', '∅')
    # Remove remaining LaTeX commands
    text = re.sub(r'\\[a-zA-Z]+', '', text)
    # Clean up extra braces
    text = text.replace('{', '').replace('}', '')
    # Clean up extra spaces
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def create_algorithm_docx(algorithms, output_path):
    """Create Word document with algorithm formatting."""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10)
    
    for i, (caption, require, ensure, lines) in enumerate(algorithms, 1):
        # Add horizontal line before algorithm
        if i > 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Add horizontal line
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '12')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '000000')
            pBdr.append(bottom)
            pPr.append(pBdr)
        
        # Algorithm title
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = title.add_run(f'Algorithm {i} {caption}')
        run.bold = True
        run.font.size = Pt(12)
        title.paragraph_format.space_after = Pt(2)
        
        # Add horizontal line after title
        pPr = title._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
        
        # Require line
        require_para = doc.add_paragraph()
        require_para.paragraph_format.space_before = Pt(4)
        require_para.paragraph_format.space_after = Pt(2)
        run_label = require_para.add_run('Require: ')
        run_label.bold = True
        run_label.font.size = Pt(10)
        run_value = require_para.add_run(require)
        run_value.italic = True
        run_value.font.size = Pt(10)
        
        # Ensure line
        ensure_para = doc.add_paragraph()
        ensure_para.paragraph_format.space_before = Pt(2)
        ensure_para.paragraph_format.space_after = Pt(4)
        run_label = ensure_para.add_run('Ensure: ')
        run_label.bold = True
        run_label.font.size = Pt(10)
        run_value = ensure_para.add_run(ensure)
        run_value.italic = True
        run_value.font.size = Pt(10)
        
        # Algorithm lines
        indent_level = 0
        for line_data in lines:
            line_type = line_data[0]
            
            if line_type == 'comment':
                comment_text = line_data[1]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run = p.add_run(f'    {comment_text}')
                run.italic = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0, 0, 0)
                
            elif line_type == 'state':
                line_num = line_data[1]
                content = line_data[2]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_content = p.add_run(content)
                run_content.font.size = Pt(10)
                
            elif line_type == 'if':
                line_num = line_data[1]
                condition = line_data[2]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_if = p.add_run('if ')
                run_if.bold = True
                run_if.font.size = Pt(10)
                run_cond = p.add_run(condition)
                run_cond.italic = True
                run_cond.font.size = Pt(10)
                run_then = p.add_run(' then')
                run_then.bold = True
                run_then.font.size = Pt(10)
                indent_level += 1
                
            elif line_type == 'else':
                indent_level -= 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run = p.add_run('else')
                run.bold = True
                run.font.size = Pt(10)
                indent_level += 1
                
            elif line_type == 'elsif':
                line_num = line_data[1]
                condition = line_data[2]
                indent_level -= 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_els = p.add_run('else if ')
                run_els.bold = True
                run_els.font.size = Pt(10)
                run_cond = p.add_run(condition)
                run_cond.italic = True
                run_cond.font.size = Pt(10)
                run_then = p.add_run(' then')
                run_then.bold = True
                run_then.font.size = Pt(10)
                indent_level += 1
                
            elif line_type == 'while':
                line_num = line_data[1]
                condition = line_data[2]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_while = p.add_run('while ')
                run_while.bold = True
                run_while.font.size = Pt(10)
                run_cond = p.add_run(condition)
                run_cond.italic = True
                run_cond.font.size = Pt(10)
                run_do = p.add_run(' do')
                run_do.bold = True
                run_do.font.size = Pt(10)
                indent_level += 1
                
            elif line_type == 'endif':
                indent_level -= 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run = p.add_run('end if')
                run.bold = True
                run.font.size = Pt(10)
                
            elif line_type == 'endwhile':
                indent_level -= 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run = p.add_run('end while')
                run.bold = True
                run.font.size = Pt(10)
                
            elif line_type == 'forall':
                line_num = line_data[1]
                condition = line_data[2]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_for = p.add_run('for all ')
                run_for.bold = True
                run_for.font.size = Pt(10)
                run_cond = p.add_run(condition)
                run_cond.italic = True
                run_cond.font.size = Pt(10)
                run_do = p.add_run(' do')
                run_do.bold = True
                run_do.font.size = Pt(10)
                indent_level += 1
                
            elif line_type == 'endfor':
                indent_level -= 1
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run = p.add_run('end for')
                run.bold = True
                run.font.size = Pt(10)
                
            elif line_type == 'return':
                line_num = line_data[1]
                content = line_data[2]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Inches(0.25 * indent_level)
                run_num = p.add_run(f'{line_num}: ')
                run_num.font.size = Pt(10)
                run_return = p.add_run('return ')
                run_return.bold = True
                run_return.font.size = Pt(10)
                if content:
                    run_content = p.add_run(content)
                    run_content.italic = True
                    run_content.font.size = Pt(10)
    
    # Add final horizontal line
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)
    
    doc.save(output_path)
    print(f"Created: {output_path}")


def main():
    psudocode_dir = Path("src/tinycua/psudocode")
    output_file = Path("TINYCUA_Algorithms.docx")
    
    algorithms = []
    
    # Process each tex file in order
    tex_files = [
        "query-analyst.tex",
        "information-digester.tex",
        "worker.tex",
        "task-processing.tex",
        "main-workflow.tex"
    ]
    
    for tex_file in tex_files:
        tex_path = psudocode_dir / tex_file
        if tex_path.exists():
            print(f"Parsing: {tex_path}")
            content = tex_path.read_text()
            parsed = parse_latex_algorithm(content)
            algorithms.append(parsed)
    
    create_algorithm_docx(algorithms, output_file)


if __name__ == "__main__":
    main()
